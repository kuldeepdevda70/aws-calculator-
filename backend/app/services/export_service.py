import io
import logging
from typing import Any, Dict, List
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from app.services.validation_service import validation_service

logger = logging.getLogger("export_service")


class ExportService:
    """
    Generates downloadable multi-tab Excel (.xlsx) workbooks and structured PDF (.pdf) reports
    covering executive KPIs, itemized workloads, inferred auxiliary architecture,
    100% resource completeness accounting, and cost reconciliation.
    """

    def generate_excel_report(self, estimate_data: Dict[str, Any]) -> io.BytesIO:
        """Creates an executive multi-tab Excel workbook."""
        wb = Workbook()

        # Styles
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        accent_fill = PatternFill(start_color="FF9900", end_color="FF9900", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        sub_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        font_title = Font(name="Calibri", size=16, bold=True, color="1E293B")
        font_section = Font(name="Calibri", size=13, bold=True, color="1E293B")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_bold = Font(name="Calibri", size=11, bold=True)
        font_regular = Font(name="Calibri", size=11)
        font_muted = Font(name="Calibri", size=10, italic=True, color="64748B")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        resources = estimate_data.get("priced_resources", [])
        inferred = estimate_data.get("inferred_infrastructure", [])
        validation = validation_service.validate_completeness(resources)
        reconciliation = validation_service.reconcile_costs(estimate_data)

        monthly_cost = float(estimate_data.get("total_monthly_cost_usd", 0.0))
        workloads_monthly = float(estimate_data.get("workloads_monthly_cost_usd", monthly_cost))
        inferred_monthly = float(estimate_data.get("inferred_monthly_cost_usd", 0.0))
        annual_cost = float(estimate_data.get("total_annual_cost_usd", monthly_cost * 12))
        region = estimate_data.get("default_region", "ap-south-1")
        total_items = estimate_data.get("total_resources_calculated", len(resources))

        # ----------------------------------------------------
        # SHEET 1: Executive Summary
        # ----------------------------------------------------
        ws_sum = wb.active
        ws_sum.title = "Executive Summary"

        ws_sum.merge_cells("A1:F1")
        ws_sum["A1"] = "AWS Cloud Migration - Financial Executive Summary"
        ws_sum["A1"].font = font_title
        ws_sum.row_dimensions[1].height = 28

        kpis = [
            ("Target AWS Region", f"{region}"),
            ("Total Migrated Workloads", total_items),
            ("Direct Workloads Cost (Monthly)", workloads_monthly),
            ("Inferred Auxiliary Infrastructure (Monthly)", inferred_monthly),
            ("Total Estimated AWS Cost (Monthly)", monthly_cost),
            ("Total Estimated AWS Cost (Annual)", annual_cost),
            ("Original Source Cloud Bill (Monthly)", float(estimate_data.get("total_source_bill_cost", 0.0))),
            ("Financial Variance Delta (USD / %)", f"${reconciliation['delta_usd']:,.2f} ({reconciliation['delta_percentage']}%)")
        ]

        ws_sum["A3"] = "Key Migration Metrics"
        ws_sum["A3"].font = font_section

        row_k = 4
        for label, val in kpis:
            ws_sum[f"A{row_k}"] = label
            ws_sum[f"A{row_k}"].font = font_bold
            ws_sum[f"B{row_k}"] = val
            ws_sum[f"B{row_k}"].font = font_regular
            if isinstance(val, (int, float)):
                ws_sum[f"B{row_k}"].number_format = '$#,##0.00'
                ws_sum[f"B{row_k}"].font = font_bold
            ws_sum[f"A{row_k}"].border = thin_border
            ws_sum[f"B{row_k}"].border = thin_border
            row_k += 1

        # Category Breakdown
        row_k += 2
        ws_sum[f"A{row_k}"] = "Category Breakdown (Monthly USD)"
        ws_sum[f"A{row_k}"].font = font_section
        row_k += 1

        ws_sum[f"A{row_k}"] = "Architecture Category"
        ws_sum[f"B{row_k}"] = "Monthly Cost (USD)"
        ws_sum[f"A{row_k}"].font = font_header
        ws_sum[f"A{row_k}"].fill = header_fill
        ws_sum[f"B{row_k}"].font = font_header
        ws_sum[f"B{row_k}"].fill = header_fill
        row_k += 1

        cat_breakdown = estimate_data.get("category_breakdown_usd", {})
        for cat_name, cat_val in cat_breakdown.items():
            ws_sum[f"A{row_k}"] = cat_name
            ws_sum[f"B{row_k}"] = float(cat_val)
            ws_sum[f"B{row_k}"].number_format = '$#,##0.00'
            ws_sum[f"A{row_k}"].border = thin_border
            ws_sum[f"B{row_k}"].border = thin_border
            row_k += 1

        ws_sum.column_dimensions["A"].width = 42
        ws_sum.column_dimensions["B"].width = 28

        # ----------------------------------------------------
        # SHEET 2: Itemized Workloads
        # ----------------------------------------------------
        ws_items = wb.create_sheet(title="Itemized Workloads")

        ws_items.merge_cells("A1:K1")
        ws_items["A1"] = "Itemized Workload Architectural Mapping & Cost Calculation"
        ws_items["A1"].font = font_title
        ws_items.row_dimensions[1].height = 28

        headers_items = [
            "Source Workload", "Source Cost", "Target AWS Service",
            "Instance / Spec", "Qty", "Storage (GB)", "Unit Rate",
            "Monthly USD", "Annual USD", "Confidence", "Mapping Rationale"
        ]

        start_row = 3
        for col_idx, h in enumerate(headers_items, 1):
            c = ws_items.cell(row=start_row, column=col_idx, value=h)
            c.font = font_header
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws_items.row_dimensions[start_row].height = 24

        curr_row = start_row + 1
        for idx, res in enumerate(resources):
            ws_items.cell(row=curr_row, column=1, value=res.get("source_service_name", ""))
            c2 = ws_items.cell(row=curr_row, column=2, value=float(res.get("source_cost", 0.0)))
            c2.number_format = '$#,##0.00'
            ws_items.cell(row=curr_row, column=3, value=res.get("target_aws_service", ""))
            ws_items.cell(row=curr_row, column=4, value=res.get("target_aws_instance_type", ""))
            ws_items.cell(row=curr_row, column=5, value=float(res.get("target_quantity", 1.0)))
            ws_items.cell(row=curr_row, column=6, value=float(res.get("target_storage_gb", 0.0)))
            c7 = ws_items.cell(row=curr_row, column=7, value=float(res.get("pricing_unit_rate", 0.0)))
            c7.number_format = '$#,##0.0000'
            c8 = ws_items.cell(row=curr_row, column=8, value=float(res.get("monthly_cost_usd", 0.0)))
            c8.number_format = '$#,##0.00'
            c9 = ws_items.cell(row=curr_row, column=9, value=float(res.get("annual_cost_usd", 0.0)))
            c9.number_format = '$#,##0.00'
            ws_items.cell(row=curr_row, column=10, value=res.get("mapping_confidence", "category_rule"))
            ws_items.cell(row=curr_row, column=11, value=res.get("mapping_rationale") or res.get("notes", ""))

            for c in range(1, 12):
                cell = ws_items.cell(row=curr_row, column=c)
                cell.font = font_regular
                cell.border = thin_border
                if idx % 2 == 1:
                    cell.fill = zebra_fill
            curr_row += 1

        # Totals row
        ws_items.cell(row=curr_row, column=1, value="TOTAL DIRECT WORKLOADS").font = font_bold
        ws_items.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=7)
        t_m = ws_items.cell(row=curr_row, column=8, value=workloads_monthly)
        t_m.font = font_bold
        t_m.number_format = '$#,##0.00'
        t_m.fill = accent_fill

        t_a = ws_items.cell(row=curr_row, column=9, value=workloads_monthly * 12)
        t_a.font = font_bold
        t_a.number_format = '$#,##0.00'
        t_a.fill = accent_fill

        col_w_items = {1: 36, 2: 14, 3: 24, 4: 20, 5: 10, 6: 14, 7: 15, 8: 16, 9: 16, 10: 16, 11: 45}
        for col_idx, width in col_w_items.items():
            ws_items.column_dimensions[chr(64 + col_idx)].width = width

        # ----------------------------------------------------
        # SHEET 3: Inferred Infrastructure
        # ----------------------------------------------------
        ws_inf = wb.create_sheet(title="Inferred Infrastructure")

        ws_inf.merge_cells("A1:G1")
        ws_inf["A1"] = "Inferred Auxiliary Architecture (Required for Production Readiness)"
        ws_inf["A1"].font = font_title
        ws_inf.row_dimensions[1].height = 28

        headers_inf = [
            "Inferred Service", "Component", "Quantity / Storage",
            "Rate (USD)", "Monthly USD", "Annual USD", "Architectural Rationale"
        ]

        for col_idx, h in enumerate(headers_inf, 1):
            c = ws_inf.cell(row=3, column=col_idx, value=h)
            c.font = font_header
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws_inf.row_dimensions[3].height = 24

        curr_row = 4
        for idx, inf_item in enumerate(inferred):
            ws_inf.cell(row=curr_row, column=1, value=inf_item.get("target_aws_service", ""))
            ws_inf.cell(row=curr_row, column=2, value=inf_item.get("target_aws_instance_type", ""))
            qty_st = inf_item.get("target_storage_gb") or inf_item.get("target_quantity", 1.0)
            ws_inf.cell(row=curr_row, column=3, value=f"{qty_st} {inf_item.get('pricing_unit_type', '')}")
            c4 = ws_inf.cell(row=curr_row, column=4, value=float(inf_item.get("pricing_unit_rate", 0.0)))
            c4.number_format = '$#,##0.0000'
            c5 = ws_inf.cell(row=curr_row, column=5, value=float(inf_item.get("monthly_cost_usd", 0.0)))
            c5.number_format = '$#,##0.00'
            c6 = ws_inf.cell(row=curr_row, column=6, value=float(inf_item.get("annual_cost_usd", 0.0)))
            c6.number_format = '$#,##0.00'
            ws_inf.cell(row=curr_row, column=7, value=inf_item.get("rationale", ""))

            for c in range(1, 8):
                cell = ws_inf.cell(row=curr_row, column=c)
                cell.font = font_regular
                cell.border = thin_border
                if idx % 2 == 1:
                    cell.fill = zebra_fill
            curr_row += 1

        ws_inf.cell(row=curr_row, column=1, value="TOTAL INFERRED INFRASTRUCTURE").font = font_bold
        ws_inf.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
        t_im = ws_inf.cell(row=curr_row, column=5, value=inferred_monthly)
        t_im.font = font_bold
        t_im.number_format = '$#,##0.00'
        t_im.fill = accent_fill
        t_ia = ws_inf.cell(row=curr_row, column=6, value=inferred_monthly * 12)
        t_ia.font = font_bold
        t_ia.number_format = '$#,##0.00'
        t_ia.fill = accent_fill

        col_w_inf = {1: 26, 2: 24, 3: 20, 4: 15, 5: 16, 6: 16, 7: 60}
        for col_idx, width in col_w_inf.items():
            ws_inf.column_dimensions[chr(64 + col_idx)].width = width

        # ----------------------------------------------------
        # SHEET 4: Completeness & Reconciliation
        # ----------------------------------------------------
        ws_val = wb.create_sheet(title="Completeness & Reconciliation")

        ws_val.merge_cells("A1:E1")
        ws_val["A1"] = "Resource Completeness Accounting & Cost Reconciliation Report"
        ws_val["A1"].font = font_title
        ws_val.row_dimensions[1].height = 28

        ws_val["A3"] = "100% Resource Completeness Accounting"
        ws_val["A3"].font = font_section

        ws_val["A4"] = "Accounting Status"
        ws_val["B4"] = "Resource Count"
        ws_val["C4"] = "Percentage (%)"
        ws_val["D4"] = "Description"
        for c in ["A4", "B4", "C4", "D4"]:
            ws_val[c].font = font_header
            ws_val[c].fill = header_fill

        val_rows = [
            ("MAPPED_VALIDATED", validation["counts"]["MAPPED_VALIDATED"], f"{validation['percentages']['MAPPED_VALIDATED']}%", "Exact SKU or spec-matched sizing"),
            ("MAPPED_ASSUMPTIONS", validation["counts"]["MAPPED_ASSUMPTIONS"], f"{validation['percentages']['MAPPED_ASSUMPTIONS']}%", "Category heuristic or default sizing assumptions"),
            ("REVIEW_REQUIRED", validation["counts"]["REVIEW_REQUIRED"], f"{validation['percentages']['REVIEW_REQUIRED']}%", "Unmapped workloads requiring architect review"),
            ("EXCLUDED", validation["counts"]["EXCLUDED"], f"{validation['percentages']['EXCLUDED']}%", "Zero-cost headers and non-infrastructure items")
        ]

        r_v = 5
        for status, cnt, pct, desc in val_rows:
            ws_val[f"A{r_v}"] = status
            ws_val[f"B{r_v}"] = cnt
            ws_val[f"C{r_v}"] = pct
            ws_val[f"D{r_v}"] = desc
            for col_l in ["A", "B", "C", "D"]:
                ws_val[f"{col_l}{r_v}"].border = thin_border
                ws_val[f"{col_l}{r_v}"].font = font_regular
            r_v += 1

        # Reconciliation section
        r_v += 2
        ws_val[f"A{r_v}"] = "Cost Reconciliation Analysis"
        ws_val[f"A{r_v}"].font = font_section
        r_v += 1

        rec_items = [
            ("Original Source Cloud Monthly Cost", reconciliation["source_monthly_bill_usd"]),
            ("AWS Direct Workloads Monthly Cost", reconciliation["aws_direct_workloads_monthly_usd"]),
            ("AWS Inferred Auxiliary Monthly Cost", reconciliation["aws_inferred_infrastructure_monthly_usd"]),
            ("Total Estimated AWS Monthly Cost", reconciliation["aws_total_estimated_monthly_usd"]),
            ("Variance Delta (USD)", reconciliation["delta_usd"]),
            ("Variance Delta (%)", f"{reconciliation['delta_percentage']}%")
        ]

        for lbl, v in rec_items:
            ws_val[f"A{r_v}"] = lbl
            ws_val[f"B{r_v}"] = v
            ws_val[f"A{r_v}"].font = font_bold
            ws_val[f"A{r_v}"].border = thin_border
            ws_val[f"B{r_v}"].border = thin_border
            if isinstance(v, (int, float)):
                ws_val[f"B{r_v}"].number_format = '$#,##0.00'
            r_v += 1

        r_v += 1
        ws_val[f"A{r_v}"] = "Primary Variance Drivers:"
        ws_val[f"A{r_v}"].font = font_bold
        r_v += 1
        for driver in reconciliation.get("variance_drivers", []):
            ws_val[f"A{r_v}"] = f"• {driver}"
            ws_val[f"A{r_v}"].font = font_muted
            r_v += 1

        ws_val.column_dimensions["A"].width = 38
        ws_val.column_dimensions["B"].width = 20
        ws_val.column_dimensions["C"].width = 18
        ws_val.column_dimensions["D"].width = 50

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def generate_pdf_report(self, estimate_data: Dict[str, Any]) -> io.BytesIO:
        """Generates a styled landscape PDF cost estimate summary table using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )

        styles = getSampleStyleSheet()
        elements = []

        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1E293B")
        )
        subtitle_style = ParagraphStyle(
            name="SubtitleStyle",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#64748B")
        )

        elements.append(Paragraph("<b>AWS Cloud Migration Cost Estimate & Architecture Plan</b>", title_style))
        region = estimate_data.get("default_region", "ap-south-1")
        monthly = float(estimate_data.get("total_monthly_cost_usd", 0.0))
        workloads_m = float(estimate_data.get("workloads_monthly_cost_usd", monthly))
        inferred_m = float(estimate_data.get("inferred_monthly_cost_usd", 0.0))
        annual = float(estimate_data.get("total_annual_cost_usd", monthly * 12))

        elements.append(Paragraph(
            f"Target Region: <b>{region}</b> | "
            f"Direct Workloads: <b>${workloads_m:,.2f}</b> | "
            f"Inferred Auxiliary: <b>${inferred_m:,.2f}</b> | "
            f"Total Monthly: <b>${monthly:,.2f} USD</b> | "
            f"Annual AWS: <b>${annual:,.2f} USD</b>",
            subtitle_style
        ))
        elements.append(Spacer(1, 15))

        # Data Table
        table_data = [[
            Paragraph("<b>Source Workload</b>", styles["Normal"]),
            Paragraph("<b>Source Cost</b>", styles["Normal"]),
            Paragraph("<b>Target AWS Service</b>", styles["Normal"]),
            Paragraph("<b>Instance / Spec</b>", styles["Normal"]),
            Paragraph("<b>Qty</b>", styles["Normal"]),
            Paragraph("<b>Monthly USD</b>", styles["Normal"]),
            Paragraph("<b>Annual USD</b>", styles["Normal"])
        ]]

        resources = estimate_data.get("priced_resources", [])[:50]
        for r in resources:
            src_name = (r.get("source_service_name") or "")[:35]
            src_cost = float(r.get("source_cost") or 0.0)
            aws_svc = r.get("target_aws_service", "")
            inst = r.get("target_aws_instance_type", "")
            qty = r.get("target_quantity", 1.0)
            m_cost = float(r.get("monthly_cost_usd") or 0.0)
            a_cost = float(r.get("annual_cost_usd") or 0.0)

            table_data.append([
                Paragraph(src_name, styles["Normal"]),
                f"${src_cost:,.2f}",
                Paragraph(aws_svc, styles["Normal"]),
                Paragraph(inst, styles["Normal"]),
                str(qty),
                f"${m_cost:,.2f}",
                f"${a_cost:,.2f}"
            ])

        t = Table(table_data, colWidths=[180, 75, 130, 110, 45, 95, 95])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (4, 1), (-1, -1), 'RIGHT'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('TOPPADDING', (0, 0), (-1, 0), 6),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
        ]))
        elements.append(t)

        # Inferred section in PDF if present
        inferred = estimate_data.get("inferred_infrastructure", [])
        if inferred:
            elements.append(Spacer(1, 15))
            elements.append(Paragraph("<b>Inferred Auxiliary Infrastructure (Production Readiness)</b>", subtitle_style))
            elements.append(Spacer(1, 6))

            inf_data = [[
                Paragraph("<b>AWS Auxiliary Service</b>", styles["Normal"]),
                Paragraph("<b>Component / Type</b>", styles["Normal"]),
                Paragraph("<b>Monthly USD</b>", styles["Normal"]),
                Paragraph("<b>Architectural Rationale</b>", styles["Normal"])
            ]]
            for inf in inferred:
                inf_data.append([
                    Paragraph(inf.get("target_aws_service", ""), styles["Normal"]),
                    Paragraph(inf.get("target_aws_instance_type", ""), styles["Normal"]),
                    f"${float(inf.get('monthly_cost_usd', 0.0)):,.2f}",
                    Paragraph(inf.get("rationale", "")[:90], styles["Normal"])
                ])
            t_inf = Table(inf_data, colWidths=[160, 140, 90, 340])
            t_inf.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#334155")),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4)
            ]))
            elements.append(t_inf)

        doc.build(elements)
        buffer.seek(0)
        return buffer


export_service = ExportService()
